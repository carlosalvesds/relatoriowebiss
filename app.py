import html
import io
import re
import zipfile
import base64
import hashlib
from tema import aplicar_tema
import importlib
from pathlib import Path
from datetime import datetime
from decimal import Decimal, InvalidOperation
import xml.etree.ElementTree as ET

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import analise_xml

# Instâncias já abertas podem conservar o módulo anterior à aba de tributos.
if not all(hasattr(analise_xml, nome) for nome in ("COLUNAS_TRIBUTOS", "tributos_nfse")) or "Alíquota ISS" not in analise_xml.COLUNAS_TRIBUTOS:
    importlib.reload(analise_xml)
from analise_xml import COLUNAS_TRIBUTOS, indicadores_nfse, tributos_nfse, aliquota_iss_nfse


def texto(elemento, caminho, padrao="N/A"):
    encontrado = elemento.find(caminho)
    if encontrado is None or not encontrado.text:
        return padrao
    return encontrado.text.strip()


def numero(valor):
    if valor in (None, "", "N/A"):
        return 0.0
    valor = str(valor).strip().replace(" ", "")
    if "," in valor:
        valor = valor.replace(".", "").replace(",", ".")
    try:
        return float(Decimal(valor))
    except (InvalidOperation, ValueError):
        return 0.0


def sem_namespace(elemento):
    """Remove namespaces para aceitar XMLs ABRASF de diferentes versões."""
    for item in elemento.iter():
        if "}" in item.tag:
            item.tag = item.tag.split("}", 1)[1]
    return elemento


def processar_xml(xml_content, tributos=None, indicadores=None, arquivo="XML"):
    try:
        root = ET.parse(xml_content).getroot()
    except (ET.ParseError, OSError, ValueError):
        return [], None

    root = sem_namespace(root)
    dados = []
    emitente = None
    for nfse in root.iter("InfNfse"):
        indicador = indicadores_nfse(nfse, arquivo)
        if indicadores is not None:
            indicadores.append(indicador)
        if tributos is not None:
            tributos.append(tributos_nfse(nfse, indicador))
        if emitente is None:
            emitente = {
                "cnpj": texto(nfse, ".//PrestadorServico/IdentificacaoPrestador/CpfCnpj/Cnpj", "")
                        or texto(nfse, ".//Prestador/CpfCnpj/Cnpj"),
                "razao_social": texto(nfse, ".//PrestadorServico/RazaoSocial", "")
                                or texto(nfse, ".//Prestador/RazaoSocial", "")
                                or texto(nfse, ".//PrestadorServico/NomeFantasia", "Empresa não informada no XML"),
            }

        cpf = texto(nfse, ".//IdentificacaoTomador/CpfCnpj/Cpf", "")
        cnpj = texto(nfse, ".//IdentificacaoTomador/CpfCnpj/Cnpj", "")
        data_original = texto(nfse, ".//DataEmissao", "")
        try:
            data_emissao = datetime.fromisoformat(data_original.replace("Z", "+00:00")).strftime("%d/%m/%Y")
        except ValueError:
            data_emissao = data_original or "N/A"

        iss_retido = texto(nfse, ".//Servico/IssRetido", "2")
        dados.append({
            "CPF/CNPJ Tomador": cpf or cnpj or "N/A",
            "Razão Social Tomador": texto(nfse, ".//Tomador/RazaoSocial"),
            "Número NFS-e": texto(nfse, ".//Numero"),
            "Valor do Serviço": numero(texto(nfse, ".//Servico/Valores/ValorServicos", "0")),
            "Alíquota ISS": aliquota_iss_nfse(nfse),
            "ISS": numero(texto(nfse, ".//ValoresNfse/ValorIss", "0")),
            "ISS Retido": "Sim" if iss_retido == "1" else "Não",
            "Data de Emissão": data_emissao,
            "Item": texto(nfse, ".//Servico/ItemListaServico"),
            "Código NBS": texto(nfse, ".//Servico/CodigoNbs"),
            "Código CNAE": texto(nfse, ".//Servico/CodigoCnae"),
            "Base de Cálculo IBSCBS": numero(texto(nfse, ".//IBSCBS/valores/vBC", "0")),
            "pIBSUF": numero(texto(nfse, ".//IBSCBS/valores/uf/pIBSUF", "0")),
            "pRedAliqUF": numero(texto(nfse, ".//IBSCBS/valores/uf/pRedAliqUF", "0")),
            "pAliqEfetUF": numero(texto(nfse, ".//IBSCBS/valores/uf/pAliqEfetUF", "0")),
            "pRedAliqMun": numero(texto(nfse, ".//IBSCBS/valores/mun/pRedAliqMun", "0")),
            "pCBS": numero(texto(nfse, ".//IBSCBS/valores/fed/pCBS", "0")),
            "pRedAliqCBS": numero(texto(nfse, ".//IBSCBS/valores/fed/pRedAliqCBS", "0")),
            "pAliqEfetCBS": numero(texto(nfse, ".//IBSCBS/valores/fed/pAliqEfetCBS", "0")),
            "vIBSUF": numero(texto(nfse, ".//IBSCBS/totCIBS/gIBS/gIBSUFTot/vIBSUF", "0")),
            "vCBS": numero(texto(nfse, ".//IBSCBS/totCIBS/gCBS/vCBS", "0")),
            "Descrição do Serviço": texto(nfse, ".//Servico/Discriminacao"),
        })
    return dados, emitente


def extrair_dados_nfse(arquivo):
    dados, emitente = [], None
    tributos, indicadores = [], []
    nome = arquivo.name.lower()
    try:
        if nome.endswith(".zip"):
            with zipfile.ZipFile(arquivo) as pacote:
                for item in pacote.namelist():
                    if item.lower().endswith(".xml") and not item.startswith("__MACOSX/"):
                        with pacote.open(item) as xml:
                            encontrados, prestador = processar_xml(io.BytesIO(xml.read()), tributos, indicadores, item)
                        dados.extend(encontrados)
                        emitente = emitente or prestador
        else:
            encontrados, emitente = processar_xml(arquivo, tributos, indicadores, arquivo.name)
            dados.extend(encontrados)
    except (zipfile.BadZipFile, RuntimeError, OSError) as erro:
        raise ValueError("O arquivo ZIP está inválido ou não pôde ser lido.") from erro
    df = pd.DataFrame(dados)
    df.attrs["tributos"] = tributos
    df.attrs["indicadores"] = indicadores
    return df, emitente


def formatar_documento(valor):
    digitos = "".join(filter(str.isdigit, str(valor)))
    if len(digitos) == 11:
        return f"{digitos[:3]}.{digitos[3:6]}.{digitos[6:9]}-{digitos[9:]}"
    if len(digitos) == 14:
        return f"{digitos[:2]}.{digitos[2:5]}.{digitos[5:8]}/{digitos[8:12]}-{digitos[12:]}"
    return str(valor or "N/A")


def moeda(valor):
    return f"R$ {numero(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def gerar_excel(df):
    df = df.copy()
    df["CPF/CNPJ Tomador"] = df["CPF/CNPJ Tomador"].map(formatar_documento)
    saida = io.BytesIO()
    monetarias = {"Valor do Serviço", "ISS", "Base de Cálculo IBSCBS", "vIBSUF", "vCBS"}
    with pd.ExcelWriter(saida, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Dados NFS-e")
        planilha = writer.sheets["Dados NFS-e"]
        for celula in planilha[1]:
            celula.fill = PatternFill("solid", fgColor="1F77B4")
            celula.font = Font(bold=True, color="FFFFFF")
            celula.alignment = Alignment(horizontal="center")
        for indice, coluna in enumerate(df.columns, 1):
            letra = get_column_letter(indice)
            for celula in planilha[letra][1:]:
                celula.alignment = Alignment(horizontal="center", vertical="center")
                if coluna == "Alíquota ISS":
                    celula.number_format = "0.00%"
                elif coluna in monetarias:
                    celula.number_format = 'R$ #,##0.00'
            maior = max((len(str(c.value or "")) for c in planilha[letra]), default=10)
            planilha.column_dimensions[letra].width = min(maior + 2, 50)
        detalhes = pd.DataFrame(df.attrs.get("tributos", []), columns=COLUNAS_TRIBUTOS)
        if len(detalhes) > 1048575:
            raise ValueError("O lote excede o limite de uma aba do Excel. Divida o lote em arquivos menores.")
        detalhes["CPF/CNPJ Tomador"] = detalhes["CPF/CNPJ Tomador"].map(formatar_documento)
        detalhes.to_excel(writer, index=False, sheet_name="Tributos por NFS-e", na_rep="Não informado")
        analise = writer.sheets["Tributos por NFS-e"]
        analise.freeze_panes = "D2"
        analise.auto_filter.ref = analise.dimensions
        analise.sheet_view.showGridLines = False
        analise.sheet_properties.tabColor = "00B2E2"
        larguras = [24] + [22] * (len(COLUNAS_TRIBUTOS) - 1)
        for i, largura in enumerate(larguras, 1):
            analise.column_dimensions[get_column_letter(i)].width = largura
        for celula in analise[1]:
            celula.fill = PatternFill("solid", fgColor="0149AF")
            celula.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
            celula.alignment = Alignment(horizontal="center", vertical="center")
        analise.row_dimensions[1].height = 28
        fonte_dados = Font(name="Arial", size=10, color="142E4D")
        fundos = [PatternFill("solid", fgColor="F0F7FF"), PatternFill("solid", fgColor="FFFFFF")]
        alinhamento = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for linha in analise.iter_rows(min_row=2):
            for celula in linha:
                celula.font = fonte_dados
                celula.fill = fundos[celula.row % 2]
                celula.alignment = alinhamento
                if isinstance(celula.value, str):
                    celula.data_type = "s"
                    celula.number_format = "@"
                elif celula.column == 2:
                    celula.number_format = "dd/mm/yyyy"
                elif COLUNAS_TRIBUTOS[celula.column - 1] == "Alíquota ISS":
                    celula.number_format = "0.00%"
                elif celula.column >= 4:
                    celula.number_format = '"R$" #,##0.00'
    return saida.getvalue()


st.set_page_config(page_title="Relatório NFS-e", page_icon="📄", layout="wide")
paleta = aplicar_tema()

logo = Path(__file__).parent / "kit-visual-diroma/assets/diroma-logo.png"
marca = base64.b64encode(logo.read_bytes()).decode() if logo.exists() else ""
st.markdown(f'<div class="portal-header"><img src="data:image/png;base64,{marca}" alt="diRoma hotéis & parques"><div><h1>Relatório NFS-e</h1><p>Análise de notas fiscais e estrutura XML</p></div></div>', unsafe_allow_html=True)
arquivo = st.file_uploader("Escolha um arquivo XML ou ZIP de NFS-e", type=["xml", "zip"])

if arquivo is not None:
    try:
        arquivo.seek(0)
        conteudo = arquivo.read()
        assinatura = ("tributos-v6-nome-emitente", arquivo.name, hashlib.sha256(conteudo).hexdigest())
        if st.session_state.get("lote_assinatura") != assinatura:
            entrada = io.BytesIO(conteudo)
            entrada.name = arquivo.name
            df, emitente = extrair_dados_nfse(entrada)
            st.session_state["lote_dados"] = (df, emitente)
            st.session_state["lote_assinatura"] = assinatura
            st.session_state.pop("lote_excel", None)
        df, emitente = st.session_state["lote_dados"]
    except ValueError as erro:
        st.error(str(erro))
        st.stop()

    if df.empty:
        st.warning("Nenhuma NFS-e foi encontrada. Verifique se o XML contém a tag InfNfse.")
        st.stop()

    if emitente:
        st.subheader(f"🏢 {emitente['razao_social']}")
        st.write(f"**CNPJ:** {formatar_documento(emitente['cnpj'])}")

    indicadores = pd.DataFrame(df.attrs["indicadores"])
    tributos = pd.DataFrame(df.attrs["tributos"], columns=COLUNAS_TRIBUTOS)
    proprios = {campo: pd.to_numeric(tributos[campo], errors="coerce").gt(0)
                for campo in ("PIS Próprio", "COFINS Próprio")}
    st.subheader("Notas com tributos informados e retenções")
    contagens = {
        "Total de notas fiscais": len(indicadores),
        "IBS informado": int(indicadores["IBS"].sum()),
        "CBS informado": int(indicadores["CBS"].sum()),
        "PIS Próprio": int(proprios["PIS Próprio"].sum()),
        "COFINS Próprio": int(proprios["COFINS Próprio"].sum()),
        "PIS retido": int(indicadores["PIS retido"].eq("Sim").sum()),
        "COFINS retido": int(indicadores["COFINS retido"].eq("Sim").sum()),
        "IRRF com valor": int(indicadores["IRRF"].sum()),
    }
    cards = list(contagens.items())
    for inicio in range(0, len(cards), 4):
        for coluna, (rotulo, valor) in zip(st.columns(4), cards[inicio:inicio + 4]):
            coluna.metric(rotulo, valor)
    aliquotas_iss = sorted(df["Alíquota ISS"].dropna().unique())
    resumo_aliquotas = " / ".join(f"{valor * 100:.2f}%".replace(".", ",") for valor in aliquotas_iss) or "Não informado"
    totais = {
        "Total IBS (UF)": df["vIBSUF"].sum(),
        "Total CBS": df["vCBS"].sum(),
        "Total de Serviços": df["Valor do Serviço"].sum(),
        "Total ISS": df["ISS"].sum(),
    }
    st.subheader("📊 Resumo dos valores")
    with st.container(key="resumo-valores"):
        for coluna, (rotulo, valor) in zip(st.columns(4), totais.items()):
            coluna.metric(rotulo, moeda(valor))
        st.metric("Alíquota ISS", resumo_aliquotas)

    exibicao = df.copy()
    exibicao["CPF/CNPJ Tomador"] = exibicao["CPF/CNPJ Tomador"].map(formatar_documento)
    monetarias = ["Valor do Serviço", "ISS", "Base de Cálculo IBSCBS", "vIBSUF", "vCBS"]
    percentuais = ["pIBSUF", "pRedAliqUF", "pAliqEfetUF", "pRedAliqMun", "pCBS", "pRedAliqCBS", "pAliqEfetCBS"]
    formatos = {coluna: moeda for coluna in monetarias}
    formatos.update({coluna: "{:.2f}" for coluna in percentuais})
    formatos["Alíquota ISS"] = lambda valor: "Não informado" if pd.isna(valor) else f"{valor * 100:.2f}%".replace(".", ",")
    st.subheader("Planilha das notas")
    with st.container(key="descricao-excel"):
        st.caption("O Excel contém os dados da tabela abaixo e a aba Tributos por NFS-e, com número, data e valores dos tributos de todas as notas.")
    coluna_download, coluna_imprimir = st.columns(2)
    download = coluna_download.empty()
    if "lote_excel" not in st.session_state:
        download.download_button("Preparando planilha Excel…", b"", "dados_nfse.xlsx", disabled=True, use_container_width=True)
        try:
            with st.spinner("Preparando o Excel com os dados das notas e os tributos. Aguarde…"):
                st.session_state["lote_excel"] = gerar_excel(exibicao)
        except (ValueError, OSError) as erro:
            download.empty()
            st.error(f"Não foi possível preparar a planilha: {erro}")
    if "lote_excel" in st.session_state:
        download.download_button("Baixar planilha Excel (.xlsx)", st.session_state["lote_excel"], "dados_nfse.xlsx",
                                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 key="baixar_excel", on_click="ignore", type="primary", use_container_width=True)
    st.dataframe(exibicao.style.format(formatos).set_properties(**{
        "background-color": paleta["portal-surface"], "color": paleta["portal-text"]
    }), use_container_width=True, hide_index=True, row_height=28)

    html_tabela = exibicao.to_html(index=False, escape=True)
    nome_emitente = html.escape(emitente["razao_social"] if emitente else "Relatório")
    relatorio_html = f"""<!doctype html><html><head><meta charset="utf-8"><title>Relatório NFS-e</title>
    <style>body{{font-family:Arial;margin:24px}}h1{{color:#1f77b4}}table{{border-collapse:collapse;width:100%;font-size:10px}}
    th{{background:#1f77b4;color:white}}th,td{{border:1px solid #ddd;padding:6px;text-align:center}}</style></head>
    <body><h1>Relatório de NFS-e</h1><h2>{nome_emitente}</h2>{html_tabela}<script>window.onload=()=>window.print()</script></body></html>"""

    if coluna_imprimir.button("🖨️ Imprimir / salvar relatório em PDF", use_container_width=True):
        components.html(relatorio_html, height=1)
