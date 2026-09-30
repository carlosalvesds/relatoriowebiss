import ast
import io
from pathlib import Path
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import load_workbook
from analise_xml import COLUNAS_TRIBUTOS, indicadores_nfse, tributos_nfse


def carregar_funcoes():
    tree = ast.parse(Path(__file__).with_name("app.py").read_text(encoding="utf-8"))
    tree.body = [n for n in tree.body if isinstance(n, (ast.Import, ast.ImportFrom, ast.FunctionDef))]
    namespace = {}
    exec(compile(tree, "app.py", "exec"), namespace)
    return namespace


class AnaliseTests(unittest.TestCase):
    def test_modulo_antigo_em_memoria(self):
        import analise_xml
        from streamlit.testing.v1 import AppTest
        colunas = analise_xml.COLUNAS_TRIBUTOS
        funcao = analise_xml.tributos_nfse
        try:
            del analise_xml.COLUNAS_TRIBUTOS
            del analise_xml.tributos_nfse
            app = AppTest.from_file(str(Path(__file__).with_name('app.py'))).run(timeout=30)
            self.assertFalse(app.exception)
            self.assertEqual(analise_xml.COLUNAS_TRIBUTOS, colunas)
            self.assertTrue(callable(analise_xml.tributos_nfse))
        finally:
            analise_xml.COLUNAS_TRIBUTOS = colunas
            analise_xml.tributos_nfse = funcao

    def test_painel_e_download(self):
        from streamlit.testing.v1 import AppTest
        source = io.BytesIO(b'<CompNfse><Nfse><InfNfse Id="n1"><Numero>1</Numero><Tomador><IdentificacaoTomador><CpfCnpj><Cpf>01234567890</Cpf></CpfCnpj></IdentificacaoTomador></Tomador><ValorPis>10</ValorPis><tpRetPisCofins>5</tpRetPisCofins><pCBS>0</pCBS><ValorIr>2</ValorIr></InfNfse></Nfse></CompNfse>')
        source.name = 'teste.xml'
        source = io.BytesIO(source.getvalue().replace(b'<ValorPis>10</ValorPis>', b'<ValorPis>10</ValorPis><ValorCofins>20</ValorCofins>'))
        source.name = 'teste.xml'
        with patch('streamlit.file_uploader', return_value=source):
            app = AppTest.from_file(str(Path(__file__).with_name('app.py'))).run(timeout=30)
            self.assertFalse(app.exception)
            self.assertEqual([m.value for m in app.metric[:8]], ['1', '0', '1', '0', '1', '1', '0', '1'])
            botao = app.get('download_button')[0]
            self.assertEqual(botao.label, 'Baixar planilha Excel (.xlsx)')
            self.assertFalse(botao.disabled)
            excel = app.session_state['lote_excel']
            self.assertEqual(app.session_state['aparencia'], 'Claro')
            self.assertEqual(app.button(key='alternar_tema').label, '🌙 Modo escuro')
            app.button(key='alternar_tema').click().run(timeout=30)
            self.assertFalse(app.exception)
            self.assertEqual(app.session_state['aparencia'], 'Escuro')
            self.assertEqual(app.button(key='alternar_tema').label, '☀️ Modo claro')
            self.assertEqual(app.session_state['lote_excel'], excel)
            self.assertTrue(any('--portal-bg:#101923' in item.value for item in app.markdown))
            app.button(key='alternar_tema').click().run(timeout=30)
            self.assertFalse(app.exception)
            self.assertTrue(any('--portal-bg:#f0f7ff' in item.value for item in app.markdown))
            workbook = load_workbook(io.BytesIO(excel))
            self.assertEqual(workbook.sheetnames, ['Dados NFS-e', 'Tributos por NFS-e'])
            self.assertEqual(len(app.dataframe), 1)
            self.assertFalse(app.selectbox)
            self.assertFalse(app.expander)
            self.assertEqual(app.dataframe[0].value.iloc[0]['CPF/CNPJ Tomador'], '012.345.678-90')
            self.assertEqual(workbook['Tributos por NFS-e']['K2'].value, 2)
            self.assertEqual(workbook['Tributos por NFS-e']['C2'].value, '012.345.678-90')
            self.assertEqual(workbook['Tributos por NFS-e']['C2'].data_type, 's')
            self.assertEqual(workbook['Dados NFS-e']['A2'].value, '012.345.678-90')
            self.assertEqual(workbook['Dados NFS-e']['C2'].value, '1')
            novo = io.BytesIO(source.getvalue().replace(b'<Numero>1</Numero>', b'<Numero>2</Numero>'))
            novo.name = source.name
            with patch('streamlit.file_uploader', return_value=novo):
                app.run(timeout=30)
            self.assertFalse(app.exception)
            atualizado = load_workbook(io.BytesIO(app.session_state['lote_excel']))
            self.assertEqual(atualizado['Dados NFS-e']['C2'].value, '2')

    def test_retencoes(self):
        esperado = [(False, False), (True, True), (False, False), (True, True),
                    (True, True), (True, False), (False, True), (False, True),
                    (False, False), (True, False)]
        for codigo, (pis, cofins) in enumerate(esperado):
            root = ET.fromstring(f'<InfNfse><ValorPis>20</ValorPis><ValorCofins>30</ValorCofins><tpRetPisCofins>{codigo}</tpRetPisCofins></InfNfse>')
            result = indicadores_nfse(root, "teste.xml")
            self.assertEqual(result["PIS retido"] == "Sim", pis)
            self.assertEqual(result["COFINS retido"] == "Sim", cofins)
            valores = tributos_nfse(root, result)
            self.assertEqual(valores['PIS Retido'], 20 if pis else 0)
            self.assertEqual(valores['PIS Próprio'], 0 if pis else 20)
            self.assertEqual(valores['COFINS Retido'], 30 if cofins else 0)
            self.assertEqual(valores['COFINS Próprio'], 0 if cofins else 30)
        result = indicadores_nfse(ET.fromstring('<InfNfse><ValorPis>20</ValorPis><pCBS>0</pCBS><ValorIr>0</ValorIr></InfNfse>'), "t.xml")
        self.assertEqual(result["PIS retido"], "Não informado")
        self.assertTrue(result["CBS"])
        self.assertFalse(result["IRRF"])
        root = ET.fromstring('<InfNfse><ValorPis>20</ValorPis></InfNfse>')
        valores = tributos_nfse(root, indicadores_nfse(root, 't.xml'))
        self.assertIsNone(valores['PIS Próprio'])
        self.assertIsNone(valores['PIS Retido'])

    def test_tributos_e_exportacao(self):
        funcs = carregar_funcoes()
        xml = '<InfNfse xmlns="urn:test" Id="001"><Numero>0007</Numero><Tomador><IdentificacaoTomador><CpfCnpj><Cnpj>00123456000199</Cnpj></CpfCnpj></IdentificacaoTomador></Tomador><DataEmissao>2026-09-29T13:00:00-03:00</DataEmissao><vIBSTot>3</vIBSTot><vIBSUF>2</vIBSUF><vIBSMun>1</vIBSMun><vCBS>9</vCBS><ValorPis>20</ValorPis><ValorCofins>30</ValorCofins><tpRetPisCofins>5</tpRetPisCofins><ValoresNfse><Aliquota>3</Aliquota><ValorIss>50</ValorIss></ValoresNfse></InfNfse>'
        tributos, indicadores = [], []
        dados, _ = funcs["processar_xml"](io.BytesIO(xml.encode()), tributos, indicadores)
        self.assertEqual(len(dados), 1)
        df = pd.DataFrame(dados)
        df.attrs['tributos'] = tributos
        book = load_workbook(io.BytesIO(funcs['gerar_excel'](df)))
        self.assertEqual(book.sheetnames, ['Dados NFS-e', 'Tributos por NFS-e'])
        self.assertEqual([c.value for c in book.worksheets[0][1]], list(df.columns))
        sheet = book.worksheets[1]
        self.assertEqual([c.value for c in sheet[1]], COLUNAS_TRIBUTOS)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['A2'].value, '0007')
        self.assertEqual(sheet['C2'].value, '00.123.456/0001-99')
        self.assertEqual(sheet['C2'].data_type, 's')
        self.assertEqual(sheet['B2'].value.strftime('%d/%m/%Y'), '29/09/2026')
        self.assertEqual([c.value for c in sheet[2]][3:], [3, 9, 0, 30, 20, 0, 50, 0, 0.03])
        self.assertEqual(sheet['D2'].data_type, 'n')
        self.assertEqual(sheet.freeze_panes, 'D2')
        self.assertEqual(sheet['L2'].number_format, '0.00%')
        self.assertEqual(book.worksheets[0]['E2'].value, 0.03)
        self.assertEqual(book.worksheets[0]['E2'].number_format, '0.00%')
        self.assertEqual(book.worksheets[0]['A2'].value, '00.123.456/0001-99')
        for row in sheet:
            for cell in row:
                self.assertEqual(cell.alignment.horizontal, 'center')
                self.assertEqual(cell.alignment.vertical, 'center')

    def test_lote_real(self):
        funcs = carregar_funcoes()
        arquivo = Path(__file__).parents[1] / 'LoteXMLNotasFiscais.zip'
        if not arquivo.exists():
            self.skipTest('Lote local indisponível')
        with arquivo.open('rb') as source:
            df, emitente = funcs['extrair_dados_nfse'](source)
        self.assertEqual(emitente['cnpj'], '00071556000683')
        self.assertEqual(emitente['razao_social'], 'ROMA HOTÉIS E REALIZAÇÕES LTDA')
        self.assertFalse(df.empty)
        self.assertEqual(df['Alíquota ISS'].unique().tolist(), [0.03])
        self.assertEqual(len(df), len(df.attrs['indicadores']))
        book = load_workbook(io.BytesIO(funcs['gerar_excel'](df)), read_only=True)
        self.assertEqual(book['Dados NFS-e'].max_row, len(df) + 1)
        self.assertEqual(book['Tributos por NFS-e'].max_row, len(df) + 1)
        self.assertEqual(book['Tributos por NFS-e'].max_column, 12)
        print('Lote validado:', len(df), 'notas; 12 colunas de tributos')


if __name__ == '__main__':
    unittest.main()
